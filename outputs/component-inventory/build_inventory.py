import json, re, pathlib, collections, html
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = pathlib.Path(__file__).resolve().parents[2]
OUT = pathlib.Path(__file__).resolve().parent
ROWS = []
purpose = {
'react':'前端页面与交互组件基础框架','react-dom':'React 页面挂载与 DOM 渲染','react-router-dom':'页面路由与导航','zustand':'画布、工作流和界面状态管理',
'paper':'二维画布、矢量图形绘制与编辑','reactflow':'流程画布、节点连接与工作流编排','@reactflow/node-resizer':'流程节点尺寸调整',
'three':'三维模型加载、材质、灯光与场景渲染','@react-three/fiber':'在 React 中组织 Three.js 三维场景','@react-three/drei':'三维场景控制器与辅助组件',
'three-stdlib':'Three.js 扩展加载器与控制工具','three-gpu-pathtracer':'三维节点 GPU 路径追踪渲染','three-mesh-bvh':'三维网格 BVH 加速与相交检测','xatlas-web':'三维网格 UV 展开工具',
'ogl':'WebGL 图形与视觉效果','ag-psd':'Photoshop PSD 文件解析与生成','@webav/av-cliper':'浏览器音视频剪辑、合成与 MP4 导出',
'@imgly/background-removal':'浏览器图片背景移除；后端不可用时的降级处理','@imgly/background-removal-node':'服务端图片背景移除与工作线程处理',
'onnxruntime-web':'浏览器端 ONNX 模型推理','onnxruntime-node':'服务端 ONNX 模型推理','sharp':'服务端图像缩放、转换与处理',
'@google/genai':'Google Gemini 多模态 AI 接口调用','@google/generative-ai':'Google Gemini 旧版 SDK 接口调用',
'class-variance-authority':'界面组件样式变体管理','clsx':'条件 CSS 类名组合','tailwind-merge':'Tailwind CSS 类名冲突合并','lucide-react':'界面图标',
'date-fns':'日期格式化和日期计算','react-day-picker':'日期选择界面','echarts':'管理与统计数据图表','element-china-area-data':'中国省市区数据与地区选择',
'i18next':'多语言资源与翻译管理','i18next-browser-languagedetector':'浏览器语言检测','react-i18next':'React 多语言界面集成',
'react-markdown':'Markdown 内容渲染','remark-gfm':'Markdown 表格、任务列表等 GFM 语法支持','qrcode.react':'前端二维码展示','uuid':'唯一标识生成',
'socket.io-client':'前端实时协作与消息通信','socket.io':'服务端实时协作与消息通信','ws':'WebSocket 通信',
'@alicloud/green20220302':'阿里云内容安全审核接口','@alicloud/openapi-client':'阿里云 OpenAPI 客户端基础支持','@alicloud/sms-sdk':'阿里云短信验证码与通知','@alicloud/tea-util':'阿里云 SDK 请求与数据处理辅助',
'@aws-sdk/client-s3':'S3 兼容对象存储访问','@aws-sdk/lib-storage':'对象存储分片与大文件上传','@aws-sdk/s3-request-presigner':'对象存储预签名访问地址生成',
'ali-oss':'阿里云 OSS 文件与素材存储','cos-nodejs-sdk-v5':'腾讯云 COS 对象存储访问','alipay-sdk':'支付宝支付接口对接','wechatpay-node-v3':'微信支付 V3 接口对接',
'@nestjs/common':'后端控制器、服务与依赖注入基础','@nestjs/core':'后端应用生命周期与核心容器','@nestjs/config':'后端环境变量与配置管理','@nestjs/jwt':'JWT 身份令牌签发与校验','@nestjs/passport':'认证策略与 NestJS 集成',
'@nestjs/platform-express':'NestJS Express HTTP 适配支持','@nestjs/platform-fastify':'NestJS Fastify HTTP 适配','@nestjs/platform-socket.io':'NestJS Socket.IO 实时通信适配','@nestjs/schedule':'定时任务调度','@nestjs/swagger':'后端 OpenAPI 接口文档','@nestjs/websockets':'后端 WebSocket 网关',
'fastify':'后端 HTTP 服务框架','@fastify/compress':'HTTP 响应压缩','@fastify/cookie':'Cookie 解析与写入','@fastify/cors':'跨域请求处理','@fastify/helmet':'HTTP 安全响应头','@fastify/multipart':'文件上传与 multipart 请求处理',
'@prisma/client':'数据库类型化访问与持久化','prisma':'数据库模型与迁移、客户端生成工具','axios':'HTTP API 请求','bcryptjs':'用户密码哈希与验证','class-transformer':'请求与响应对象转换','class-validator':'请求参数与 DTO 校验','cookie':'Cookie 数据解析与序列化','dotenv':'环境变量文件加载','exceljs':'Excel 数据导入导出','ioredis':'Redis 连接、缓存及共享状态访问','json-rules-engine':'业务规则条件计算','jsonwebtoken':'JWT 令牌处理','passport':'身份认证策略框架','passport-jwt':'JWT 认证策略','qrcode':'服务端支付等二维码生成','reflect-metadata':'装饰器类型元数据支持','rxjs':'异步数据流与 NestJS 响应处理','socks-proxy-agent':'外部服务请求的 SOCKS 代理连接','multer':'Express multipart 文件上传支持',
'typescript':'TypeScript 类型检查与编译','vite':'前端开发服务器与生产构建','@vitejs/plugin-react':'Vite React 编译与热更新','tailwindcss':'前端原子化 CSS 样式生成','tailwindcss-animate':'界面动画样式扩展','postcss':'CSS 构建处理','autoprefixer':'CSS 浏览器前缀生成','eslint':'代码静态检查','@eslint/js':'JavaScript 静态检查规则','eslint-plugin-react-hooks':'React Hooks 使用规则检查','eslint-plugin-react-refresh':'React 热更新相关静态检查','typescript-eslint':'TypeScript ESLint 解析与规则','globals':'静态检查环境全局变量定义','ts-node':'开发及维护脚本运行 TypeScript','ts-node-dev':'后端开发运行与自动重启','undici':'HTTP 客户端及相关开发验证','vitest':'自动化测试执行',
}

def usage(n):
    if n.startswith('@types/'): return 'TypeScript 类型声明；用于开发与编译，不作为业务运行代码'
    if n.startswith('@radix-ui/'): return '前端无障碍 UI 组件：'+n.removeprefix('@radix-ui/react-')
    if n.startswith('@tiptap/'): return '富文本编辑器及图片、链接、占位符等编辑扩展'
    if n.startswith('@opentelemetry/'): return '服务端调用链追踪、自动埋点与 OTLP 遥测上报'
    if n.startswith('@img/sharp-'): return 'Sharp 平台可选图像处理二进制；仅对应部署平台使用'
    return purpose.get(n,'项目开发工具支持')

def provider(n, meta):
    groups=[('@nestjs/','NestJS 项目'),('@fastify/','Fastify 开源社区'),('@radix-ui/','Radix UI 项目'),('@tiptap/','Tiptap 开源项目'),('@google/','Google'),('@imgly/','IMG.LY GmbH'),('@aws-sdk/','Amazon Web Services'),('@alicloud/','阿里云'),('@opentelemetry/','OpenTelemetry 项目'),('@prisma/','Prisma'),('@react-three/','Poimandres 开源社区'),('@types/','DefinitelyTyped 开源社区'),('@img/','Sharp 项目'),('@webav/','WebAV 开源项目')]
    for prefix, name in groups:
        if n.startswith(prefix): return name
    known={'react':'React 开源社区（Meta 发起）','react-dom':'React 开源社区（Meta 发起）','three':'Three.js 开源社区','paper':'Paper.js 项目','prisma':'Prisma','ali-oss':'阿里云','alipay-sdk':'支付宝开放平台','cos-nodejs-sdk-v5':'腾讯云','wechatpay-node-v3':'wechatpay-node-v3 开源项目（非官方 SDK）','onnxruntime-web':'Microsoft ONNX Runtime 项目','onnxruntime-node':'Microsoft ONNX Runtime 项目','echarts':'Apache ECharts 项目','typescript':'Microsoft','xatlas-web':'Mozilla Reality / xatlas-web 项目','reactflow':'xyflow / React Flow 项目','@reactflow/node-resizer':'xyflow / React Flow 项目'}
    if n in known:return known[n]
    a=meta.get('author')
    if isinstance(a,dict): a=a.get('name')
    if isinstance(a,str):return re.sub(r'\s*<[^>]+>','',a).strip()
    return n+' 开源项目'

for directory in ['.','frontend','backend']:
    base=ROOT/directory
    pkg=json.loads((base/'package.json').read_text(encoding='utf-8'))
    lock=json.loads((base/'package-lock.json').read_text(encoding='utf-8'))
    files=[]
    # Scan source and build/config files, never environment files or user uploads.
    for folder in ['src','scripts','tests']:
        if (base/folder).exists() and directory!='.':
            files += [f for f in (base/folder).rglob('*') if f.suffix in {'.ts','.tsx','.js','.mjs','.cjs','.css'} and f.is_file()]
    files += [f for f in base.glob('*') if f.is_file() and f.suffix in {'.ts','.js','.mjs','.cjs'}]
    texts=[(str(f.relative_to(ROOT)).replace('\\','/'),f.read_text(encoding='utf-8',errors='replace')) for f in files]
    for section in ['dependencies','optionalDependencies','devDependencies']:
        for n, declared in pkg.get(section,{}).items():
            entry=lock['packages'].get('node_modules/'+n,{})
            meta={}
            try:meta=json.loads((base/'node_modules'/n/'package.json').read_text(encoding='utf-8'))
            except (OSError,ValueError):pass
            lic=entry.get('license') or (meta.get('license') if meta.get('version')==entry.get('version') else None) or '未标注，待核验'
            license_source='package-lock.json / '+n
            if n.startswith('@imgly/background-removal'):
                lic='AGPL v3（LICENSE.md）'
                license_source=str(base.relative_to(ROOT))+'/node_modules/'+n+'/LICENSE.md'
            if n=='@webav/av-cliper':
                lic='上游 MIT；1.2.8 包未声明，版本许可待核验'
                license_source='https://github.com/WebAV-Tech/WebAV/blob/main/LICENSE'
            refs=[path for path,txt in texts if re.search(r'''['"]'''+re.escape(n)+r'''(?:/[^'"\n]*)?['"]''',txt)]
            runtime=section!='devDependencies' and not n.startswith('@types/')
            state='检出代码引用' if refs else ('开发/构建配置依赖' if not runtime else '依赖清单声明；未检出直接代码引用')
            scope=usage(n)
            if directory=='.':scope='根目录依赖声明；'+scope+'；未检出根目录业务入口引用'
            elif not refs and runtime:scope+='；当前仅确认依赖声明'
            if not runtime and not n.startswith('@types/'): scope+='；开发/构建使用'
            ROWS.append(dict(module='根目录' if directory=='.' else directory,name=n,version=entry.get('version','未锁定'),declared=declared,license=lic,provider=provider(n,meta),scope=scope,kind='运行/可选依赖' if runtime else '开发/类型依赖',evidence=refs,reference_status=state,license_source=license_source))

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.5);sec.right_margin=Cm(1.5)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10)
normal._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.line_spacing=1.12
for sty in ['Title','Heading 1','Heading 2']:
    doc.styles[sty].font.color.rgb=RGBColor(0,0,0)
    doc.styles[sty]._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体')
p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('表 6-3  开源及第三方组件清单');r.bold=True;r.font.size=Pt(16)
p.paragraph_format.space_after=Pt(10)
p=doc.add_paragraph('范围：Tanva 根目录、前端及后端，不包含 tgAgent。版本采用当前 npm 锁文件中的精确版本；统计日期为 2026 年 10 月 16 日。列出直接依赖，开发工具另表列示，不等同于全部传递依赖或生产部署软件物料清单。')
p.paragraph_format.space_after=Pt(8)

def table(headers, records, widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'6');b.set(qn('w:color'),'555555');borders.append(b)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,value in [('top','85'),('bottom','85'),('left','95'),('right','95')]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    header=t.rows[0]
    repeat=OxmlElement('w:tblHeader');header._tr.get_or_add_trPr().append(repeat)
    for c,txt in zip(header.cells,headers):
        c.text=txt;s=OxmlElement('w:shd');s.set(qn('w:fill'),'DCE6EE');c._tc.get_or_add_tcPr().append(s)
        for r in c.paragraphs[0].runs:r.bold=True
    for record in records:
        row=t.add_row()
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for cell,txt,w in zip(row.cells,record,widths):cell.text=txt;cell.width=Cm(w)
    for row in t.rows:
        for c,w in zip(row.cells,widths):
            c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph().paragraph_format.space_after=Pt(3)
    return t

for kind,title in [('运行/可选依赖','业务运行及可选组件'),('开发/类型依赖','开发构建及类型组件')]:
    doc.add_heading(title,level=1)
    records=[]
    for row in ROWS:
        if row['kind']!=kind:continue
        records.append([row['name']+'\n'+row['version']+'（'+row['module']+'）',row['provider'],row['license'],row['scope']])
    table(['名称及版本','来源/提供方','许可协议/授权方式','在本作品中的用途'],records,[5.1,3.5,3.7,5.7])

doc.add_heading('部署及外部资源',level=1)
infra=[
['Redis\n7-alpine（浮动标签）','Redis 项目','具体版本待确认；≤7.2 为 BSD-3-Clause；7.4 系列为 RSALv2 或 SSPLv1','后端缓存与共享状态；Compose 配置启用。不得据此认定实际运行版本。'],
['OpenObserve\nv0.70.2（镜像标签）','OpenObserve 项目','上游社区版 AGPL-3.0；镜像对应授权待核验','日志及链路观测平台；Compose 配置启用。'],
['PostgreSQL\n部署版本未锁定','PostgreSQL Global Development Group','PostgreSQL License','后端关系数据库；Prisma 数据源为 postgresql。Compose 中 16-alpine 示例已注释。'],
['Nginx\n版本未锁定','Nginx 开源项目','BSD-2-Clause（开源版）','前端静态文件服务与 API 反向代理；仓库存在 nginx.conf。'],
['glTF 示例模型\nsimple-box.glb / duck.glb','文档注明来自 Khronos glTF-Sample-Assets','应按具体模型文件附带许可确认，当前资源目录未附具体许可','前端三维模型加载与渲染示例。模型资源不应套用 JavaScript 库许可证。'],
]
table(['名称及版本','来源/提供方','许可协议/授权方式','在本作品中的用途'],infra,[5.1,3.5,3.7,5.7])
doc.add_heading('说明与核验依据',level=1)
notes=[
'1. 本表按组件在各目录的声明分别列示。同名包不同版本保留独立记录；版本取 package-lock.json，不以 package.json 的 ^ 或 ~ 范围代替精确版本。backend/pnpm-lock.yaml 未作为本表版本基准。',
'2. 许可字段优先取锁文件；必要时核对同版本已安装包的 package.json 与 LICENSE.md。IMG.LY 前后端抠图库的本地许可文本均为 GNU Affero General Public License v3；本表不表示已取得商业授权。',
'3. WebAV 1.2.8 已安装包未标注 license，锁文件为 ???；上游当前仓库为 MIT，尚不能据此代替该版本的完整授权证据。Sharp wasm 可选包许可组合包含 LGPL-3.0-or-later，已按锁文件单列。',
'4. “当前仅确认依赖声明”表示未在本次扫描范围内检出直接包引用，不认定为已启用功能；构建工具、动态加载、间接引用及运行条件可能影响实际使用。商业云服务的 SDK 开源许可与云服务合同、模型及数据许可分别适用。',
'5. 代码引用范围：frontend 与 backend 的 src、scripts、tests 及顶层 JS/TS 构建配置；不读取环境变量文件。引用明细与机器可读记录保存在同目录 component-evidence.json。',
'6. 来源文件：package.json / package-lock.json、frontend/package.json / package-lock.json、backend/package.json / package-lock.json、backend/docker-compose.yml、backend/prisma/schema.prisma、frontend/nginx.conf、frontend/public/models/README.md。',
'7. 上游核验链接：IMG.LY https://github.com/imgly/background-removal-js；WebAV https://github.com/WebAV-Tech/WebAV/blob/main/LICENSE；Redis https://redis.io/legal/licenses/；OpenObserve https://github.com/openobserve/openobserve。',
]
for txt in notes:
    p=doc.add_paragraph(txt);p.paragraph_format.space_after=Pt(6)
doc.core_properties.title='Tanva 开源及第三方组件清单';doc.core_properties.author=''
doc.save(OUT/'Tanva-开源及第三方组件清单.docx')
(OUT/'Tanva-开源及第三方组件清单.html').write_text('''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>Tanva 开源及第三方组件清单</title><style>
*{box-sizing:border-box}body{margin:32px auto;max-width:1120px;padding:0 24px;color:#111;background:white;font-family:"SimSun","Songti SC",serif;font-size:17px;line-height:1.5}h1{text-align:center;font-size:24px;margin:0 0 22px;font-family:"SimHei",sans-serif}h2{font-size:20px;margin:28px 0 12px;color:#000}p{margin:10px 0 18px}table{border-collapse:collapse;width:100%;table-layout:fixed;margin:0 0 20px}th,td{border:1px solid #555;padding:10px 12px;vertical-align:middle;text-align:left;overflow-wrap:anywhere}th{background:#dce6ee;font-family:"SimHei",sans-serif;font-size:18px}td:first-child{width:28%}thead{display:table-header-group}tr{break-inside:avoid}small{font-size:15px}.notes{font-size:15px}button{font:16px sans-serif;padding:8px 16px;cursor:pointer;margin-bottom:20px}@media print{body{margin:0;max-width:none;padding:0;font-size:10pt}h1{font-size:16pt}h2{font-size:13pt;break-after:avoid}th{font-size:10.5pt}td,th{padding:6px 7px}button{display:none}.notes{font-size:9pt}@page{size:A4 portrait;margin:16mm 13mm}}
</style><body><button onclick="window.print()">打印或另存为 PDF</button><h1>表 6-3&nbsp; 开源及第三方组件清单</h1><p>范围：Tanva 根目录、前端及后端，不包含 tgAgent。版本采用当前 npm 锁文件中的精确版本；统计日期为 2026 年 10 月 16 日。下表列出直接依赖，开发工具另表列示，不等同于全部传递依赖或生产部署软件物料清单。</p>'''+''.join('<h2>'+title+'</h2><table><colgroup><col style="width:28%"><col style="width:19%"><col style="width:21%"><col style="width:32%"></colgroup><thead><tr>'+''.join('<th>'+s+'</th>' for s in ['名称及版本','来源/提供方','许可协议/授权方式','在本作品中的用途'])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+html.escape(s).replace('\n','<br>')+'</td>' for s in record)+'</tr>' for record in records)+'</tbody></table>' for title,records in [(title,[[r['name']+'\n'+r['version']+'（'+r['module']+'）',r['provider'],r['license'],r['scope']] for r in ROWS if r['kind']==kind]) for kind,title in [('运行/可选依赖','业务运行及可选组件'),('开发/类型依赖','开发构建及类型组件')]]+[('部署及外部资源',infra)])+'<h2>说明与核验依据</h2><div class="notes">'+''.join('<p>'+html.escape(n)+'</p>' for n in notes)+'</div></body></html>',encoding='utf-8')
(OUT/'component-evidence.json').write_text(json.dumps({'date':'2026-10-16','scope':['root','frontend','backend'],'excluded':['tgAgent'],'version_basis':'npm package-lock.json','direct_dependencies':ROWS,'deployment_and_assets':infra},ensure_ascii=False,indent=2),encoding='utf-8')
print('Direct dependency records:',len(ROWS));print(collections.Counter(r['kind'] for r in ROWS));print('DOCX:',OUT/'Tanva-开源及第三方组件清单.docx')
